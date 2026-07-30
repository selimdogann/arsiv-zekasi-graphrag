"""
PdfLoader, DocxLoader ve AutoDocumentLoader için testler.

DocxLoader gerçek bir .docx dosyasıyla test edilir (python-docx hem yazıp
hem okuyabildiği için ucuz ve gerçekçi). PdfLoader ise pypdf.PdfReader'ı
SAHTE bir nesneyle değiştirerek (monkeypatch) test edilir — amacımız
pypdf'in kendi ayrıştırmasını test etmek değil, KENDİ kodumuzun (sayfaları
birleştirme, hata çevirme) doğru çalıştığını kanıtlamak.
"""
import docx
import pytest

from graphrag.domain.exceptions import IngestionError
from graphrag.infrastructure.ingestion.auto_loader import AutoDocumentLoader
from graphrag.infrastructure.ingestion.docx_loader import DocxLoader
from graphrag.infrastructure.ingestion.pdf_loader import PdfLoader


def test_pdf_loader_sayfalari_birlestirir(monkeypatch, tmp_path):
    class _SahtePage:
        def __init__(self, metin):
            self._metin = metin

        def extract_text(self):
            return self._metin

    class _SahteReader:
        def __init__(self, uri):
            self.pages = [_SahtePage("birinci sayfa\n"), _SahtePage("ikinci sayfa")]

    monkeypatch.setattr(
        "graphrag.infrastructure.ingestion.pdf_loader.PdfReader", _SahteReader
    )

    dosya = tmp_path / "sozde.pdf"
    dosya.write_bytes(b"%PDF-1.4 sahte icerik")

    document = PdfLoader().load(str(dosya))

    assert document.raw_text == "birinci sayfa\nikinci sayfa"
    assert document.title == "sozde"


def test_pdf_loader_dosya_yoksa_ingestionerror_firlatir():
    with pytest.raises(IngestionError):
        PdfLoader().load("/olmayan/dosya.pdf")


def test_docx_loader_paragraflari_birlestirir(tmp_path):
    dosya = tmp_path / "sozlesme.docx"
    belge = docx.Document()
    belge.add_paragraph("Birinci paragraf")
    belge.add_paragraph("İkinci paragraf")
    belge.save(str(dosya))

    document = DocxLoader().load(str(dosya))

    assert document.raw_text == "Birinci paragraf\nİkinci paragraf"
    assert document.title == "sozlesme"


def test_docx_loader_dosya_yoksa_ingestionerror_firlatir():
    with pytest.raises(IngestionError):
        DocxLoader().load("/olmayan/dosya.docx")


def test_auto_loader_txt_uzantisini_dogru_yonlendirir(tmp_path):
    dosya = tmp_path / "metin.txt"
    dosya.write_text("düz metin içerik", encoding="utf-8")

    document = AutoDocumentLoader().load(str(dosya))

    assert document.raw_text == "düz metin içerik"


def test_auto_loader_docx_uzantisini_dogru_yonlendirir(tmp_path):
    dosya = tmp_path / "belge.docx"
    belge = docx.Document()
    belge.add_paragraph("word içeriği")
    belge.save(str(dosya))

    document = AutoDocumentLoader().load(str(dosya))

    assert document.raw_text == "word içeriği"


def test_auto_loader_desteklenmeyen_uzanti_hata_firlatir(tmp_path):
    dosya = tmp_path / "resim.png"
    dosya.write_bytes(b"sahte png")

    with pytest.raises(IngestionError):
        AutoDocumentLoader().load(str(dosya))
