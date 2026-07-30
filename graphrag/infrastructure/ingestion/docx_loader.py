"""
DocxLoader — `IDocumentLoader` sözleşmesinin Word (.docx) uygulaması.

Bir .docx dosyasını paragraf paragraf okuyup metnini birleştirir, bir
`Document` nesnesine çevirir. `PlainTextLoader` ile birebir aynı iskelet —
sadece dosya okuma kısmı farklı.
"""
from __future__ import annotations

import os

import docx

from graphrag.domain.entities import Document
from graphrag.domain.exceptions import IngestionError
from graphrag.domain.interfaces import IDocumentLoader


class DocxLoader(IDocumentLoader):
    def load(self, uri: str) -> Document:
        if not os.path.isfile(uri):
            raise IngestionError(f"Dosya bulunamadı: {uri}")
        try:
            belge = docx.Document(uri)
            text = "\n".join(paragraf.text for paragraf in belge.paragraphs)
        except Exception as exc:
            raise IngestionError(f"Word dosyası okunamadı: {uri}") from exc

        title = os.path.splitext(os.path.basename(uri))[0]
        return Document(uri=uri, title=title, raw_text=text)
